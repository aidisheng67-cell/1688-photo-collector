from openpyxl import Workbook, load_workbook

from sheets import FeishuSheet, XlsxSheet, _Base, _col_letter

COLS = {"id": "订单号", "image_url": "图片链接", "confirm": "确认", "status": "状态"}


def test_col_letter():
    assert _col_letter(0) == "A"
    assert _col_letter(25) == "Z"
    assert _col_letter(26) == "AA"
    assert _col_letter(51) == "AZ"


def test_is_confirm():
    base = _Base({}, ["是", "确认", "1", "true", "yes"])
    assert base._is_confirm(" 是 ")
    assert base._is_confirm("TRUE")
    assert not base._is_confirm("否")
    assert not base._is_confirm(None)


def test_feishu_txt_variants():
    assert FeishuSheet._txt(None) == ""
    assert FeishuSheet._txt(True) == "true"
    assert FeishuSheet._txt(12) == "12"
    assert FeishuSheet._txt([{"text": "a"}, {"text": "b"}]) == "ab"
    assert FeishuSheet._txt({"text": "x"}) == "x"


def test_xlsx_read_and_write(tmp_path):
    path = tmp_path / "sheet.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["订单号", "图片链接", "确认", "状态"])
    ws.append(["1001", "https://x/a.jpg", "是", ""])
    ws.append(["1002", "https://x/b.jpg", "否", "已完成"])
    wb.save(path)

    sheet = XlsxSheet({"path": str(path)}, COLS, ["是", "确认", "1", "true", "yes"])
    rows = sheet.read_rows()

    assert [r["id"] for r in rows] == ["1001", "1002"]
    assert rows[0]["confirm"] is True
    assert rows[1]["confirm"] is False
    assert rows[1]["status"] == "已完成"
    assert rows[0]["_rownum"] == 2

    sheet.write_back(2, {"status": "处理中"})
    wb2 = load_workbook(path)
    assert wb2.active.cell(2, 4).value == "处理中"
    wb2.close()
