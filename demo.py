"""离线演示：不需要任何 API 凭证，生成示例表格并跑一轮监听。

运行：python demo.py
演示内容：xlsx 表格后端、确认列去抖/触发、状态回写、结构化日志。
真实采集需要 onebound 凭据，见 README。
"""
import os
import shutil

from openpyxl import Workbook, load_workbook

from watcher import Watcher

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_demo")


def make_sheet(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "订单"
    ws.append(["订单号", "图片链接", "确认", "状态", "保存路径", "处理时间"])
    ws.append(["DEMO-1001", "https://example.com/a.jpg", "是", "", "", ""])
    ws.append(["DEMO-1002", "https://example.com/b.jpg", "否", "", "", ""])
    ws.append(["DEMO-1003", "https://example.com/c.jpg", "是", "已完成", "旧路径", "2026-01-01 00:00:00"])
    wb.save(path)
    wb.close()


def main():
    shutil.rmtree(BASE, ignore_errors=True)
    os.makedirs(os.path.join(BASE, "samples"), exist_ok=True)
    sheet_path = os.path.join(BASE, "samples", "orders.xlsx")
    make_sheet(sheet_path)

    cfg = {
        "paths": {
            "work": os.path.join(BASE, "work"),
            "output": os.path.join(BASE, "output"),
            "cache": os.path.join(BASE, "cache", "cache.sqlite"),
            "log": os.path.join(BASE, "logs"),
            "review": os.path.join(BASE, "review"),
        },
        "sheet": {
            "backend": "xlsx",
            "columns": {
                "id": "订单号",
                "image_url": "图片链接",
                "confirm": "确认",
                "status": "状态",
                "save_path": "保存路径",
                "time": "处理时间",
            },
            "confirm_values": ["是", "确认", "1", "true", "yes"],
            "confirm_delay_seconds": 0,
            "poll_seconds": 30,
            "stale_minutes": 30,
            "xlsx": {"path": sheet_path},
        },
        "matcher": {
            "match_high": 5, "match_review": 10,
            "orb_high": 20, "orb_weak": 10,
            "orb_nfeatures": 4000, "orb_ratio": 0.8,
        },
        "thresholds": {"dedup": 5, "img_dedup": 5},
        "junk_filter": {},
        "collect": {"max_sellers": 6, "max_scan": 25, "max_detail": 30, "rerun_extra": 3},
    }

    watcher = Watcher(cfg, dry_run=True)
    processed = watcher.tick()
    watcher.obs.report()

    wb = load_workbook(sheet_path)
    ws = wb.active
    print("\n回写后的表格：")
    for row in ws.iter_rows(min_row=1, values_only=True):
        print("  " + " | ".join("" if v is None else str(v) for v in row))
    wb.close()
    print(f"\n本轮触发 {processed} 单；DEMO-1002 未确认，保持不动；DEMO-1003 已完成，直接跳过。")
    print(f"运行产物在 {BASE}")


if __name__ == "__main__":
    main()
